"""Browser test: a computer WITHOUT Tesseract and WITHOUT LibreOffice.

    needs:  a second server started like this (port 5056):
              FEA_DISABLE_OCR=1 FEA_DISABLE_LIBREOFFICE=1 python app/server.py --no-browser --port 5056
    run:    python tests/browser_ui_manual_mode.py
"""
import glob, os, tempfile
from playwright.sync_api import sync_playwright
from pptx import Presentation
import pymupdf
from pathlib import Path as _P
_HERE = _P(__file__).resolve().parent
_ROOT = _HERE.parent
_SHOTS = _HERE / "_shots"; _SHOTS.mkdir(exist_ok=True)
import shutil, tempfile
import os as _os
URL = _os.environ.get("FEA_URL", "http://127.0.0.1:5056"); SH = str(_SHOTS) + "/"
NAMED = _P(tempfile.mkdtemp(prefix="fea_named_"))              # the 9 example pictures under file names that tell what they are
for src, dst in {"image-2": "case1_bc", "image-1": "case1_deformation", "image-3": "case1_stress", "image-5": "case2_bc", "image-4": "case2_deformation",
                 "image-6": "case2_stress", "image-7": "geometry_top", "image-8": "geometry_bottom", "image-9": "mesh"}.items():
    shutil.copyfile(_ROOT / "examples" / (src + ".png"), NAMED / (dst + ".png"))
log, dialogs, fails = [], [], []
def check(name, cond, info=""):
    print(("  ok   " if cond else "  FAIL ") + name + (f"   [{info}]" if (info and not cond) else ""))
    if not cond: fails.append(name)

import os
LAUNCH = {"executable_path": os.environ.get("CHROMIUM")} if os.environ.get("CHROMIUM") else {}
if LAUNCH: LAUNCH["args"] = ["--no-sandbox"]
with sync_playwright() as p:
    b = p.chromium.launch(**LAUNCH)
    ctx = b.new_context(viewport={"width": 1280, "height": 900}, accept_downloads=True)
    page = ctx.new_page()
    page.on("console", lambda m: log.append((m.type, m.text))); page.on("pageerror", lambda e: log.append(("pageerror", str(e))))
    page.on("dialog", lambda d: (dialogs.append(d.message), d.dismiss()))
    page.goto(URL); page.wait_for_selector(".chip")
    print("chips:", page.locator("#chips").inner_text().replace("\n", " | "))
    check("chip says manual mode", "manual mode" in page.locator("#chips").inner_text())
    check("chip says the PDF comes from slide pictures", "slide pictures" in page.locator("#chips").inner_text())
    al = page.locator("#alert").inner_text()
    check("a yellow heads-up explains it - and says you can carry on", "can carry on" in al and "Tesseract" in al, al[:150])
    check("the upload box is NOT disabled", page.locator("#drop").is_visible())

    print("1. named pictures, no OCR")
    files = sorted(glob.glob(str(NAMED / "*.png")))
    page.set_input_files("#file", files)
    page.wait_for_selector("#s-review:not([hidden])", timeout=60000)
    page.wait_for_function("S.build.status === 'ok'", timeout=30000)
    check("the review page opened (no error, no 503)", True)
    check("intro says 'Manual mode'", "Manual mode" in page.locator(".card.intro").inner_text())
    roles = page.evaluate("S.images.map(i => [i.name, i.role, i.case])")
    check("roles from the file names: 2 cases", sum(1 for r in roles if r[1] == "bc") == 2 and sum(1 for r in roles if r[1] == "stress") == 2, str(roles))
    page.wait_for_function("document.querySelector('#pv-img').naturalWidth > 100", timeout=20000)
    check("the first slide picture is shown", True)
    n0 = page.evaluate("S.build.slides.length")
    print(f"     preview before typing any number: {n0} slides; notes: {page.locator('#pv-notes').inner_text()[:160]!r}")
    check("a preview exists although no number is known yet", n0 >= 7)
    page.screenshot(path=SH + "p5b_manual_review.png")

    print("2. the engineer types the numbers (what OCR would have read)")
    for cn, (d, s_) in {1: ("0.085", "74.2"), 2: ("2.42", "545")}.items():
        page.locator(f"[data-casecard='{cn}'] input[data-label='Case {cn} max deformation']").fill(d)
        page.locator(f"[data-casecard='{cn}'] input[data-label='Case {cn} max stress']").fill(s_)
    page.locator("input[data-yield='0']").fill("380")
    page.wait_for_function(f"S.build.status === 'ok' && S.build.slides.length > {n0}", timeout=30000)
    n1 = page.evaluate("S.build.slides.length")
    check("observation slides appear once the numbers are typed", n1 > n0, f"{n0} -> {n1}")
    obs2 = " ".join(page.evaluate("S.cases[1].obs"))
    check("Case 2 (545 MPa > yield 380) is reported as above yield", "above the yield strength" in obs2, obs2[:200])

    print("3. downloads on a computer without LibreOffice")
    with page.expect_download(timeout=60000) as dl:
        page.locator("#btn-pptx").click()
    d = dl.value; tmp = tempfile.mkdtemp(); f1 = os.path.join(tmp, d.suggested_filename); d.save_as(f1)
    prs = Presentation(f1)
    check("PowerPoint downloaded and opens", len(prs.slides) == n1, d.suggested_filename)
    alltext = " ".join(sh.text_frame.text for sl in prs.slides for sh in sl.shapes if sh.has_text_frame)
    check("the typed numbers are in the report", "0.085" in alltext and "74.2" in alltext and "2.42" in alltext and "545" in alltext)
    with page.expect_download(timeout=90000) as dl2:
        page.locator("#btn-pdf").click()
    d2 = dl2.value; f2 = os.path.join(tmp, d2.suggested_filename); d2.save_as(f2)
    check("PDF downloaded (made from slide pictures) with one page per slide", len(pymupdf.open(f2)) == n1, d2.suggested_filename)
    texts = page.locator(".toast").all_inner_texts()
    check("the person is told the PDF text cannot be selected", any("cannot be selected" in t for t in texts), str(texts)[:200])
    check("no dialog appeared", not dialogs, str(dialogs))

    print("4. unnamed pictures and no OCR: nothing can be guessed, still no dead end")
    page.get_by_role("button", name="Use other pictures").click()
    page.wait_for_selector("#s-upload:not([hidden])")
    ex = sorted(glob.glob(str(_ROOT / "examples" / "*.png")))
    page.set_input_files("#file", ex)
    page.wait_for_selector("#s-review:not([hidden])", timeout=60000)
    page.wait_for_function("S.build.status === 'ok'", timeout=30000)
    unk = page.evaluate("S.images.filter(i => i.role === 'unknown').length")
    check("9 pictures wait for the engineer's choice", unk == 9, str(unk))
    check("the preview still shows a (short) report", page.evaluate("S.build.slides.length") >= 2, str(page.evaluate("S.build.slides.length")))
    with page.expect_download(timeout=60000) as dl3:
        page.locator("#btn-pptx").click()
    check("and it can be downloaded", dl3.value.suggested_filename.endswith(".pptx"))
    sels = page.locator("table.pics select.sel")
    sels.nth(6).select_option("geometry"); page.wait_for_timeout(2500)
    page.wait_for_function("S.build.status === 'ok'", timeout=30000)
    check("choosing a role by hand adds the slide (no OCR needed for a geometry picture)", any("Geometry" in s["title"] for s in page.evaluate("S.build.slides")), str(page.evaluate("S.build.slides.map(s => s.title)")))
    errs = [l for l in log if l[0] in ("error", "pageerror") and "503" not in l[1]]
    check("no JavaScript errors", not errs, str(errs[:3]))
    b.close()
print("\nFAILS:", fails if fails else "none")
