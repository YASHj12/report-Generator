#!/usr/bin/env python3
"""
Self-test of the server (no browser, no running server needed):   python tests/test_api.py
Covers the error paths and edge cases a real user can hit.
"""
import io
import random
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "app"))
sys.path.insert(0, str(ROOT / "engine"))
import analyzer                                  # noqa: E402
import server                                    # noqa: E402
from PIL import Image                            # noqa: E402
from pptx import Presentation                    # noqa: E402

EX = ROOT / "examples"
c = server.app.test_client()
fails = []


def check(name, cond, info=""):
    print(("  ok   " if cond else "  FAIL ") + name + (f"   [{info}]" if info and not cond else ""))
    if not cond:
        fails.append(name)


def upload(files, endpoint="/api/analyze"):
    return c.post(endpoint, data={"files": [(io.BytesIO(b), n) for n, b in files]}, content_type="multipart/form-data")


def payload_from(a, **over):
    cover = dict(a["cover"])
    cases = [{"n": k["n"], "name": k["name"], "subtitle": k["subtitle"], "short_name": k["short_name"], "bc_id": k["bc_id"],
              "def_id": k["def_id"], "stress_id": k["stress_id"],
              "bc_items": [t.replace("[location - please confirm]", "the shaft") for t in k["bc_items"]], "notes": [], "extras": [],
              "def_max": k["def_max"], "stress_max": k["stress_max"]} for k in a["cases"]]
    p = {"sid": a["sid"], "cover": cover, "short": cover["short"], "fos": 1.3, "yield": 380, "allowable": 292, "show_utilisation": False,
         "material": {"rows": [["IS 2062", "Duct", "191400", "0.3", "7.8E-9", "380", "292"]], "units": server.SETTINGS["units"]},
         "geometry": a["geometry"], "meshes": a["meshes"], "cases": cases, "make_pdf": False}
    p.update(over)
    return p


print("1. cross-site guard (only active when the program runs on this computer)")
body = {"allowable": 292, "short": "Blade", "cases": []}
check("request from another website is refused", c.post("/api/derive", json=body, headers={"Origin": "http://evil.example"}).status_code == 403)
check("request from the page itself is accepted", c.post("/api/derive", json=body, headers={"Origin": "http://localhost"}).status_code == 200)
check("request without Origin (curl, tests) is accepted", c.post("/api/derive", json=body).status_code == 200)

print("2. bad uploads")
r = upload([("notes.pdf", b"%PDF-1.4 not a picture")])
check("a PDF is refused with a plain message", r.status_code == 400 and "could be opened as pictures" in r.get_json()["error"], r.get_data(as_text=True)[:120])
r = upload([("fake.png", b"this is not a png")])
check("a broken .png is refused", r.status_code == 400, r.get_data(as_text=True)[:120])
r = c.post("/api/analyze", data={}, content_type="multipart/form-data")
check("an empty upload is refused", r.status_code == 400)

print("3. picture reader (Tesseract) missing -> manual mode, never a dead end")
orig = analyzer.find_tesseract
analyzer.find_tesseract = lambda: None
st = c.get("/api/status").get_json()
named = [("case1_bc.png", "image-2.png"), ("case1_deformation.png", "image-1.png"), ("case1_stress.png", "image-3.png"),
         ("geometry_top.png", "image-7.png"), ("mesh.png", "image-9.png")]
rm = upload([(n, (EX / src).read_bytes()) for n, src in named]); am = rm.get_json()
ru = upload([("image-1.png", (EX / "image-1.png").read_bytes())]); au = ru.get_json()
ra = c.post("/api/analyze_example")
analyzer.find_tesseract = orig
check("status reports the reader as missing", st["ocr"] is False)
check("status still offers the preview and a PDF", st["preview"] is True and st["pdf"] is True)
check("analysis answers 200 (no 503) and says 'manual'", rm.status_code == 200 and am.get("manual") is True, rm.get_data(as_text=True)[:150])
check("the install hint is a warning in the list, not an error", any("Tesseract" in w["text"] for w in am["warnings"]))
roles = {i["name"]: (i["role"], i["case"]) for i in am["images"]}
check("file names tell the roles and the case", roles["case1_bc.png"] == ("bc", 1) and roles["case1_stress.png"] == ("stress", 1)
      and roles["geometry_top.png"][0] == "geometry" and roles["mesh.png"][0] == "mesh", roles)
check("1 case with setup + deformation + stress, 1 geometry, 1 mesh", len(am["cases"]) == 1 and len(am["geometry"]) == 1 and len(am["meshes"]) == 1)
check("an unnamed picture waits for the engineer's choice (still HTTP 200)", ru.status_code == 200 and au["images"][0]["role"] == "unknown")
check("the example button works in manual mode too", ra.status_code == 200)
mp = payload_from(am)
mp["cases"][0].update(def_max="0.085", stress_max="74.2", bc_items=["Fixed support applied on the shaft."])     # the engineer types what OCR would have read
g = c.post("/api/generate", json=mp)
check("a report is built from a manual-mode analysis once the numbers are typed", g.status_code == 200 and g.get_json()["slides"] >= 8, g.get_data(as_text=True)[:200])
if g.status_code == 200:
    (server.OUT / g.get_json()["pptx"]).unlink(missing_ok=True)

print("4. JPG pictures instead of PNG")
jp = []
for p in sorted(EX.glob("*.png")):
    buf = io.BytesIO(); Image.open(p).convert("RGB").save(buf, "JPEG", quality=88); jp.append((p.stem + ".jpg", buf.getvalue()))
random.Random(2).shuffle(jp)
r = upload(jp); a = r.get_json()
roles = {i["name"]: (i["role"], i["max"]) for i in a["images"]}
check("9 JPGs analysed", r.status_code == 200 and len(a["images"]) == 9)
check("roles and Max values survive JPG compression",
      roles["image-1.jpg"] == ("deformation", 0.085) and roles["image-3.jpg"] == ("stress", 74.2) and roles["image-4.jpg"] == ("deformation", 2.42)
      and roles["image-6.jpg"] == ("stress", 545.0) and roles["image-9.jpg"][0] == "mesh" and roles["image-2.jpg"][0] == "bc", roles)

print("5. a job with one load case only")
one = [(n, (EX / n).read_bytes()) for n in ("image-7.png", "image-9.png", "image-2.png", "image-1.png", "image-3.png")]
r = upload(one); a1 = r.get_json()
check("1 case, 1 geometry view", len(a1["cases"]) == 1 and len(a1["geometry"]) == 1, [x["n"] for x in a1["cases"]])
g = c.post("/api/generate", json=payload_from(a1))
check("PowerPoint generated", g.status_code == 200, g.get_data(as_text=True)[:200])
if g.status_code == 200:
    f = server.OUT / g.get_json()["pptx"]
    titles = []
    for s in Presentation(str(f)).slides:
        t = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        titles.append(t[0] if t else "")
    check("8 slides (cover, material, geometry, mesh, setup, results, observation, summary)", len(titles) == 8, titles)
    check("single case: slide titles have no 'Case 1:' prefix on results", any(x.strip() == "Results" for x in titles), titles)
    f.unlink(missing_ok=True)

print("6. the user corrects a picture (regroup)")
r = upload(sorted(((p.name, p.read_bytes()) for p in EX.glob("*.png"))))
a = r.get_json()
sid = a["sid"]
asg = [{"id": i["id"], "role": i["role"], "case": i["case"]} for i in a["images"]]
asg[2]["role"] = "unused"; asg[2]["case"] = None
rg = c.post("/api/regroup", json={"sid": sid, "assign": asg}).get_json()
check("dropping the Case-1 stress picture is reported", any("no von-Mises" in w["text"] for w in rg["warnings"]), rg["warnings"])
asg[2]["role"] = "stress"; asg[2]["case"] = 1
rg = c.post("/api/regroup", json={"sid": sid, "assign": asg}).get_json()
check("putting it back clears the warning", not rg["warnings"], rg["warnings"])

print("7. missing things never block: the report is built anyway and says what it did")
good = payload_from(a)
def gen(**over):
    r = c.post("/api/generate", json={**good, **over}); j = r.get_json() or {}
    return r.status_code, j
def deck_of(j):
    f = server.OUT / j["pptx"]
    prs = Presentation(str(f)); titles, texts = [], []
    for sl in prs.slides:
        tt = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        titles.append(tt[0] if tt else "")
        for sh in sl.shapes:
            if sh.has_text_frame: texts.append(sh.text_frame.text)
            if sh.has_table: texts += [c_.text for r_ in sh.table.rows for c_ in r_.cells]
    f.unlink(missing_ok=True)
    return titles, " | ".join(texts)
code, j = gen(); base_titles, _ = deck_of(j)
check("reference: the complete job has 11 slides", code == 200 and len(base_titles) == 11, base_titles)
nomat = {"material": {"rows": [["IS 2062", "Duct", "191400", "0.3", "7.8E-9", "", ""]], "units": server.SETTINGS["units"]}}
code, j = gen(allowable=None, **{"yield": None}, **nomat); t, txt = deck_of(j)
check("missing yield / allowable: built, placeholders in the material table and the summary", code == 200 and txt.count("[value - please confirm]") >= 3, txt.count("[value - please confirm]"))
check("... and no pass / fail verdict was invented", "within the materials" not in txt and "exceeds the materials" not in txt and "inadequate" not in txt and "[Conclusion - please confirm" in txt)
code, j = gen(meshes=[]); t, _ = deck_of(j)
check("missing mesh picture: built without the mesh slide", code == 200 and len(t) == 10 and not any("Mesh" in x for x in t), t)
check("... and it says so", any("Mesh slide left out" in n for n in j["notes"]), j["notes"])
code, j = gen(geometry=[]); t, _ = deck_of(j)
check("missing geometry: built, and the mesh section is renumbered to 2", code == 200 and len(t) == 10 and any(x.startswith("2.") and "Mesh" in x for x in t), t)
code, j = gen(geometry=[], meshes=[]); t, _ = deck_of(j)
check("neither geometry nor mesh: built, 9 slides", code == 200 and len(t) == 9, t)
code, j = gen(geometry=[], meshes=[], missing="frame"); t, _ = deck_of(j)
check("policy 'frame': both slides stay as empty frames (11 slides)", code == 200 and len(t) == 11 and any("empty frame" in n for n in j["notes"]), (len(t), j["notes"]))
bad = [dict(good["cases"][0], bc_items=[]), good["cases"][1]]
code, j = gen(cases=bad); t, _ = deck_of(j)
check("empty boundary-condition list: built (the setup picture stays)", code == 200 and len(t) == 11, t)
bad = [dict(good["cases"][0], stress_max="abc"), good["cases"][1]]
code, j = gen(cases=bad); t, txt = deck_of(j)
check("non-numeric Max stress: built, that sentence is simply left out, the table shows a dash", code == 200 and "\u2013" in txt and "Therefore" in txt, txt[-300:])
bad = [dict(good["cases"][0], def_id=None), good["cases"][1]]
code, j = gen(cases=bad); t, _ = deck_of(j)
check("a case without its deformation picture: built, its stress plot fills the slide", code == 200 and any("Case 1" in x and "Stress" in x for x in t), t)
bad = [dict(good["cases"][0], bc_id=None, def_id=None, stress_id=None), good["cases"][1]]
code, j = gen(cases=bad); t, _ = deck_of(j)
check("a case without ANY picture: built, the numbers and sentences stay", code == 200 and any("Observation" in x for x in t), t)
code, j = gen(cases=[], geometry=[], meshes=[]); t, _ = deck_of(j)
check("no pictures and no cases at all: still a report (cover + material)", code == 200 and len(t) == 2, t)
first_img = next(i for i in server.load_infos(a["sid"]) if i["role"] == "stress")
keep = Path(first_img["path"]).read_bytes()
Path(first_img["path"]).write_bytes(b"this is not a picture any more")
code, j = gen(); t, _ = deck_of(j)
Path(first_img["path"]).write_bytes(keep)
check("a picture file that became unreadable is treated as missing, not as a crash", code == 200 and len(t) >= 9, (code, t))
code, j = gen(cover={"title": "Bad \x00\x01 characters", "report_no": "X/1", "date": "", "client": "A\x0cB"}); t, _ = deck_of(j)
check("control characters pasted into a text box cannot break the file", code == 200, j)
r = c.post("/api/generate", json={**good, "sid": "../../etc"}); check("a forged session id is refused", r.status_code == 404)
check("download cannot escape the outputs folder", c.get("/download/..%2F..%2Fsettings.json").status_code in (404, 400))

print("8. wording follows the numbers")
d = c.post("/api/derive", json={"allowable": 292, "yield": 380, "fos": 1.3, "short": "Blade", "show_utilisation": False, "cases": [
    {"n": 1, "def_max": 0.085, "stress_max": 74.2, "legend": [74.2, 66, 57.7, 49.5, 41.2], "red_px": 0},
    {"n": 2, "def_max": 2.42, "stress_max": 545, "legend": [545, 292, 256, 219, 183], "red_px": 0}]}).get_json()
t1, t2 = d["cases"]["1"], d["cases"]["2"]
check("within-limit case says adequate", t1["exceeds"] is False and "adequate" in t1["texts"]["observations"][-1] and not t1["texts"]["conclusion"])
check("exceeding case says inadequate + corrective action", t2["exceeds"] is True and "inadequate" in t2["texts"]["observations"][-1] and "corrective" in t2["texts"]["conclusion"])
d = c.post("/api/derive", json={"allowable": None, "short": "Blade", "cases": [{"n": 1, "def_max": 1.0, "stress_max": 55.0}]}).get_json()
tt = d["cases"]["1"]; ob = " ".join(tt["texts"]["observations"]).lower()
check("without a yield stress no pass/fail text is invented", tt["exceeds"] is None and not any(w in ob for w in ("within", "exceed", "adequate")), ob)
check("... the verdict is a visible placeholder and the numbers are still stated", "please confirm" in ob and "55.0 mpa" in ob and "is 1 mm" in ob, ob)
d = c.post("/api/derive", json={"allowable": 292, "yield": 380, "short": "Blade", "cases": [{"n": 1, "def_max": 1.0, "stress_max": None}]}).get_json()
ob = d["cases"]["1"]["texts"]["observations"]
check("a missing stress number leaves out its sentences (only the deformation sentence remains)", len(ob) == 1 and ob[0] == "Total deformation in the Blade is 1 mm.", ob)
d = c.post("/api/derive", json={"allowable": 292, "yield": 380, "short": "Blade", "cases": [{"n": 1, "def_max": None, "stress_max": 74.2}]}).get_json()
ob = d["cases"]["1"]["texts"]["observations"]
check("a missing deformation number leaves out its sentence (stress + verdict remain)", len(ob) == 2 and "74.2" in ob[0] and "adequate" in ob[1], ob)
d = c.post("/api/derive", json={"allowable": 0, "yield": 0, "short": "Blade", "cases": [{"n": 1, "def_max": 1.0, "stress_max": 5.0}]})
check("a zero yield stress is treated as 'not entered' (no division by zero)", d.status_code == 200 and d.get_json()["cases"]["1"]["tier"] is None, d.get_data(as_text=True)[:150])

print("9. preview and download (the browser saves the file into the Downloads folder)")
b1 = c.post("/api/build", json=good); m1 = b1.get_json()
check("build answers with the slide list", b1.status_code == 200 and len(m1["slides"]) == 11 and m1["slides"][1]["title"].startswith("1. Material"), b1.get_data(as_text=True)[:150])
check("the same state is not built twice (same key)", c.post("/api/build", json=good).get_json()["key"] == m1["key"])
other = payload_from(a, yield_=0) if False else {**good, "cover": {**good["cover"], "client": "Someone else"}}
check("a changed text gives a new build", c.post("/api/build", json=other).get_json()["key"] != m1["key"])
sl = c.get(f"/api/slide/{a['sid']}/{m1['key']}/5?w=960")
im = Image.open(io.BytesIO(sl.data))
check("a slide is delivered as a picture, drawn without LibreOffice", sl.status_code == 200 and sl.mimetype == "image/png" and im.size == (960, 540), (sl.status_code, im.size))
sm = c.get(f"/api/slide/{a['sid']}/{m1['key']}/0?w=240")
check("small thumbnails are JPEG", sm.status_code == 200 and sm.mimetype == "image/jpeg")
check("a slide number out of range / a forged key / a forged session are refused",
      c.get(f"/api/slide/{a['sid']}/{m1['key']}/99").status_code == 404 and c.get(f"/api/slide/{a['sid']}/{'0'*12}/0").status_code == 404
      and c.get(f"/api/slide/..%2F..%2Fx/{m1['key']}/0").status_code == 404)
dp = c.get(f"/download/{a['sid']}/{m1['key']}/pptx")
cd = dp.headers.get("Content-Disposition", "")
check("download is a normal browser download (attachment + file name)", dp.status_code == 200 and "attachment" in cd and "FEA_Report_MWI-FEA-ACK-6203-01.pptx" in cd, cd)
check("... and the file is a complete PowerPoint", len(Presentation(io.BytesIO(dp.data)).slides) == 11)
check("... a copy is kept in the outputs folder", (server.OUT / "FEA_Report_MWI-FEA-ACK-6203-01.pptx").exists())
check("only pptx and pdf can be downloaded", c.get(f"/download/{a['sid']}/{m1['key']}/exe").status_code == 404)

print("10. no LibreOffice on this computer: the PDF is made by the program itself")
import os
os.environ["FEA_DISABLE_LIBREOFFICE"] = "1"
try:
    st = c.get("/api/status").get_json()
    check("status: PDF and preview are still offered", st["pdf"] and st["preview"] and st["pdf_method"] == "images", st)
    pr = c.post("/api/pdf", json={"sid": a["sid"], "key": m1["key"]}).get_json()
    check("the PDF is made from the slide pictures and the person is told", pr["ok"] and pr["method"] == "images" and "cannot be selected" in pr["message"], pr)
    dd = c.get(f"/download/{a['sid']}/{m1['key']}/pdf")
    check("the PDF downloads as an attachment", dd.status_code == 200 and dd.mimetype == "application/pdf" and "attachment" in dd.headers.get("Content-Disposition", ""))
    try:
        import pymupdf
        doc = pymupdf.open(stream=dd.data, filetype="pdf")
        check("the PDF has one 16:9 page per slide", len(doc) == 11 and abs(doc[0].rect.width / doc[0].rect.height - 16 / 9) < 0.01, (len(doc), doc[0].rect))
    except ImportError:
        print("  (pymupdf is not installed here - page check skipped)")
    g = c.post("/api/generate", json={**good, "make_pdf": True}).get_json()
    check("the older /api/generate also produces a PDF without LibreOffice", g["pdf"] and (server.OUT / g["pdf"]).exists(), g)
finally:
    del os.environ["FEA_DISABLE_LIBREOFFICE"]
for f in server.OUT.glob("FEA_Report_*"):
    f.unlink(missing_ok=True)

print("11. the optional engineering evidence reaches the PowerPoint (assumptions, mesh statistics, sign-off, limits, locations, section view)")
ids = {i["name"]: i["id"] for i in a["images"]}
ev = payload_from(a, show_utilisation=True, def_limit=3,
                  assumptions=["Linear elastic, isotropic material behaviour (no plasticity).", "Design temperature 210 \u00b0C: yield stress taken at that temperature."],
                  mesh_stats={"type": "SOLID187", "elements": "654321", "nodes": "1234567", "size": "10", "skew_avg": "0.28", "skew_max": "0.85", "oq_min": "0.45"})
ev["cover"] = {**ev["cover"], "prepared_by": "M. Shah", "checked_by": "P. Nair", "ref": "KCP-6203-R2", "revision": "R2"}
ev["cases"][1].update(loc_stress="the shaft-to-plate weld", loc_def="the free end of the shaft", react_applied="1000", react_sum="1004")
ev["cases"][0]["extras"] = [{"id": ids["image-8.png"], "kind": "stress", "heading": "Section A-A through the shaft support", "caption": "Section A-A: maximum 41.7 MPa at the shaft"}]
r = c.post("/api/generate", json=ev)
check("built with every optional field filled in", r.status_code == 200, r.get_data(as_text=True)[:200])
if r.status_code == 200:
    f = server.OUT / r.get_json()["pptx"]; prs = Presentation(str(f))
    titles = []
    alltext = []
    for sl in prs.slides:
        tt = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        titles.append(tt[0] if tt else "")
        for sh in sl.shapes:
            if sh.has_text_frame: alltext.append(sh.text_frame.text)
            if sh.has_table: alltext += [c_.text for r_ in sh.table.rows for c_ in r_.cells]
    txt = "\n".join(alltext).replace("\u00a0", " ")
    check("assumptions slide with the chosen lines", "4. Assumptions & Scope" in titles and "Linear elastic" in txt and "Design temperature 210" in txt, titles)
    check("a slide titled with the engineer's words for the section view", any("Section A-A through the shaft support" in t for t in titles), titles)
    check("mesh statistics line", "Elements: 654,321" in txt and "Max. skewness 0.85" in txt)
    check("sign-off lines on the cover", "Prepared by: M. Shah" in txt and "Checked by: P. Nair" in txt and "Ref.: KCP-6203-R2" in txt)
    check("deformation limit in the sentence and as a summary column", "allowable deformation limit of 3 mm" in txt and "Allowable Deformation" in txt)
    check("utilisation of the passing case", "utilisation 25 %" in txt)
    check("yield note for the 545 MPa peak", "above the yield strength of 380" in txt)
    check("location and reaction sentences", "occurs at the shaft-to-plate weld" in txt and "occurs at the free end of the shaft" in txt and "Reaction force check" in txt)
    check("a fully filled-in job leaves no 'please confirm' placeholder in the report", "please confirm" not in txt, [l for l in alltext if "please confirm" in l][:3])
    f.unlink(missing_ok=True)

print()
print("ALL TESTS PASSED" if not fails else f"{len(fails)} FAILED: {fails}")
sys.exit(1 if fails else 0)
