"""Browser test (Playwright + Chromium): the preview appears by itself, nothing blocks, browser-style downloads, missing pictures.

    needs:  a running server    python app/server.py --no-browser      (port 5055)
    run:    python tests/browser_ui_preview_download.py
"""
import sys, time, tempfile, os, re
from playwright.sync_api import sync_playwright
from pptx import Presentation
import pymupdf
from pathlib import Path as _P
_HERE = _P(__file__).resolve().parent
_ROOT = _HERE.parent
_SHOTS = _HERE / "_shots"; _SHOTS.mkdir(exist_ok=True)
import os as _os
URL = _os.environ.get("FEA_URL", "http://127.0.0.1:5055"); SH = str(_SHOTS) + "/"
log, dialogs, fails = [], [], []
def check(name, cond, info=""):
    print(("  ok   " if cond else "  FAIL ") + name + (f"   [{info}]" if (info and not cond) else ""))
    if not cond: fails.append(name)

import os
LAUNCH = {"executable_path": os.environ.get("CHROMIUM")} if os.environ.get("CHROMIUM") else {}
if LAUNCH: LAUNCH["args"] = ["--no-sandbox"]
with sync_playwright() as p:
    b = p.chromium.launch(**LAUNCH)
    ctx = b.new_context(viewport={"width": 1280, "height": 900}, accept_downloads=True)
    page = ctx.new_page()
    page.on("console", lambda m: log.append((m.type, m.text))); page.on("pageerror", lambda e: log.append(("pageerror", str(e))))
    page.on("dialog", lambda d: (dialogs.append(d.message), d.dismiss()))
    page.goto(URL); page.wait_for_selector(".chip")
    chips = page.locator("#chips").inner_text()
    print("chips:", chips.replace("\n", " | "))

    print("1. the preview appears directly (nobody clicks 'Preview' or 'Generate')")
    page.get_by_role("button", name="Try it with the 9 example pictures").click()
    t0 = time.time()
    page.wait_for_selector("#s-review:not([hidden])", timeout=120000)
    t_review = time.time() - t0
    page.wait_for_function("document.querySelector('#pv-img') && document.querySelector('#pv-img').naturalWidth > 100", timeout=30000)
    t_img = time.time() - t0
    print(f"     review page after {t_review:.1f}s, first slide picture after {t_img:.1f}s")
    check("slide picture is shown without any click", page.evaluate("document.querySelector('#pv-img').naturalWidth") > 500)
    page.wait_for_function("document.querySelector('#pv-status-text').textContent.includes('up to date')", timeout=20000)
    check("status chip: 'Preview is up to date'", "up to date" in page.locator("#pv-status-text").inner_text())
    check("counter says Slide 1 of 11", page.locator("#pv-count").inner_text() == "Slide 1 of 11", page.locator("#pv-count").inner_text())
    check("filmstrip has 11 thumbnails", page.locator("#pv-strip button").count() == 11, str(page.locator("#pv-strip button").count()))
    check("title of slide 1 is the cover", "Static Structural" in page.locator("#pv-title").inner_text(), page.locator("#pv-title").inner_text())
    page.locator("#pv-card").scroll_into_view_if_needed(); page.wait_for_timeout(500)
    page.screenshot(path=SH + "p5_review_top.png")

    print("2. navigation: arrows, thumbnails, keyboard, full screen")
    for _ in range(3): page.locator("#pv-next").click()
    page.wait_for_timeout(300)
    check("3 x next -> Slide 4 of 11", page.locator("#pv-count").inner_text() == "Slide 4 of 11")
    page.locator("#pv-strip button").nth(5).click(); page.wait_for_timeout(300)
    check("thumbnail 6 -> Slide 6", page.locator("#pv-count").inner_text() == "Slide 6 of 11")
    page.keyboard.press("ArrowLeft"); page.wait_for_timeout(200)
    check("ArrowLeft -> Slide 5", page.locator("#pv-count").inner_text() == "Slide 5 of 11", page.locator("#pv-count").inner_text())
    page.locator("#pv-frame").click(); page.wait_for_timeout(600)
    check("full-screen viewer opens", page.locator("#slidebox").is_visible())
    page.wait_for_function("document.querySelector('#slidebox img').naturalWidth > 1000", timeout=15000)
    check("full-screen picture is high resolution", page.evaluate("document.querySelector('#slidebox img').naturalWidth") >= 1400)
    page.screenshot(path=SH + "p5_fullscreen.png")
    page.locator("#slidebox .sb-nav.next").click(); page.wait_for_timeout(400)
    check("viewer next -> caption shows slide 6", "Slide 6 of 11" in page.locator("#slidebox figcaption").inner_text())
    page.keyboard.press("Escape"); page.wait_for_timeout(200)
    check("Escape closes it", not page.locator("#slidebox").is_visible())

    print("3. THE complaint: required fields are still empty - downloading must work anyway")
    need_n = page.locator("#s-review .need").count()
    print(f"     {need_n} yellow items are still open; bar says: {page.locator('#bar-msg').inner_text()!r}")
    check("there ARE open yellow items (yield etc.)", need_n > 0)
    with page.expect_download(timeout=60000) as dl:
        page.locator("#btn-pptx").click()
    d = dl.value
    check("a browser download started (PowerPoint)", d.suggested_filename == "FEA_Report_MWI-FEA-ACK-6203-01.pptx", d.suggested_filename)
    tmp = tempfile.mkdtemp(); f1 = os.path.join(tmp, d.suggested_filename); d.save_as(f1)
    prs = Presentation(f1)
    check("the downloaded PowerPoint opens and has 11 slides", len(prs.slides) == 11, str(len(prs.slides)))
    txt = " ".join(sh.text_frame.text for sl in prs.slides for sh in sl.shapes if sh.has_text_frame)
    tbl = " ".join(c.text for sl in prs.slides for sh in sl.shapes if sh.has_table for r in sh.table.rows for c in r.cells)
    check("empty fields became visible [placeholders], not errors", "[value - please confirm]" in tbl, tbl[:120])
    page.wait_for_selector(".toast", timeout=5000)
    tt = page.locator(".toast").first.inner_text()
    print("     toast:", tt.replace("\n", " / ")[:200])
    check("toast tells where the file went (Downloads folder)", "Downloads" in tt)
    check("stepper moved to step 3 (Download)", page.locator("#steps li.on").get_attribute("data-stage") == "done")
    page.screenshot(path=SH + "p5_after_download.png")
    with page.expect_download(timeout=90000) as dl2:
        page.locator("#btn-pdf").click()
    d2 = dl2.value; f2 = os.path.join(tmp, d2.suggested_filename); d2.save_as(f2)
    pdf = pymupdf.open(f2)
    check("a browser download started (PDF) with 11 pages", d2.suggested_filename.endswith(".pdf") and len(pdf) == 11, f"{d2.suggested_filename} {len(pdf)}")
    check("no alert/confirm dialog ever appeared", not dialogs, str(dialogs))

    print("4. editing updates the preview by itself")
    key_before = page.evaluate("S.build.key")
    page.locator("input[data-yield='0']").fill("380")
    page.wait_for_function("S.build.status !== 'ok'", timeout=5000)
    check("status switched away from 'up to date' right after typing", True)      # wait_for_function above proves it
    page.wait_for_function(f"S.build.status === 'ok' && S.build.key !== '{key_before}'", timeout=20000)
    check("a new build appeared by itself after the edit", True)
    page.locator("#pv-strip button").nth(1).click(); page.wait_for_timeout(1200)
    page.screenshot(path=SH + "p5_material_slide.png")

    print("5. a missing picture never blocks: set the geometry pictures to 'Not used'")
    n_before = page.locator("#pv-strip button").count()
    sels = page.locator("table.pics select.sel")
    changed = 0
    for i in range(sels.count()):
        if sels.nth(i).input_value() == "geometry":
            sels.nth(i).select_option("unused"); page.wait_for_timeout(2500); changed += 1
            sels = page.locator("table.pics select.sel")
    print(f"     {changed} geometry pictures set to 'Not used'")
    page.wait_for_function("S.build.status === 'ok'", timeout=30000); page.wait_for_timeout(1500)
    page.wait_for_function("S.build.status === 'ok'", timeout=30000)
    n_after = page.locator("#pv-strip button").count()
    check("the geometry slide disappeared (11 -> 10 slides)", n_after == n_before - 1, f"{n_before} -> {n_after}")
    notes = page.locator("#pv-notes").inner_text()
    check("notes explain what was left out", "Geometry slide left out" in notes, notes[:100])
    titles = page.evaluate("S.build.slides.map(s => s.title)")
    check("section numbers closed up (Mesh is now '2.')", any(t.startswith("2. ") and "Mesh" in t for t in titles), str(titles[:4]))
    page.locator("#pv-card").scroll_into_view_if_needed(); page.wait_for_timeout(300)
    page.screenshot(path=SH + "p5_missing_geometry.png")
    page.locator("#pv-missing").select_option("frame")
    page.wait_for_function("S.build.status === 'ok' && S.build.slides.length === 11", timeout=30000)
    check("policy 'keep an empty frame' brings the slide back as a frame", "empty frame" in page.locator("#pv-notes").inner_text(), page.locator("#pv-notes").inner_text()[:100])
    page.locator("#pv-strip button").nth(2).click(); page.wait_for_timeout(1200)
    page.screenshot(path=SH + "p5_frame_slide.png")
    errs = [l for l in log if l[0] in ("error", "pageerror")]
    check("no JavaScript errors in the console", not errs, str(errs[:3]))
    b.close()
print("\nFAILS:", fails if fails else "none")
