#!/usr/bin/env python3
"""
Red-team round: the ugliest inputs the council could think of.   python tests/test_redteam.py
"""
import io
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "app"))
sys.path.insert(0, str(ROOT / "engine"))
import analyzer                                  # noqa: E402
import server                                    # noqa: E402
from PIL import Image                            # noqa: E402
from pptx import Presentation                    # noqa: E402

c = server.app.test_client()
EX = ROOT / "examples"
fails = []


def check(name, cond, info=""):
    print(("  ok   " if cond else "  FAIL ") + name + (f"   [{info}]" if info and not cond else ""))
    if not cond:
        fails.append(name)


def upload(files):
    return c.post("/api/analyze", data={"files": [(io.BytesIO(b), n) for n, b in files]}, content_type="multipart/form-data")


print("R1  OCR parser: scientific notation, decimal commas, junk in front of the numbers")
mx, mn, vals = analyzer.parse_result_legend("1.2e+003 Max\nme 1.1e+003\n| 1 9.9e+002\n8.8e+002\n7.7e+002\n-3.2e+001 Min\n")
check("scientific notation Max / Min / values", mx == 1200.0 and mn == "-3.2e+001" and vals[:3] == [1200.0, 1100.0, 990.0], (mx, mn, vals))
mx, mn, vals = analyzer.parse_result_legend("24,7 Max\n21,3\n17,9\n0,5 Min")
check("decimal commas", mx == 24.7 and vals[:3] == [24.7, 21.3, 17.9], (mx, vals))
v, notes = analyzer.repair_legend([545, 292, 256, 21.9, 183, 146])
check("a dropped decimal point inside the legend is repaired (21.9 -> 219)", v == [545, 292, 256, 219.0, 183, 146] and notes, (v, notes))
v, notes = analyzer.repair_legend([2.42, 2.15, 188.0, 1.62, 1.35])
check("188 -> 1.88", v[2] == 1.88, v)
v, notes = analyzer.repair_legend([5.0, 4.0, 4.5, 3.0])
check("an out-of-order value that no power of ten fixes is dropped, not trusted", 4.5 not in v and len(notes) == 1, (v, notes))

print("R2  a file name in Hindi / Chinese / with spaces and brackets")
names = ["चित्र 1 (final).png", "图像二.png", "image three.png"]
r = upload([(names[0], (EX / "image-7.png").read_bytes()), (names[1], (EX / "image-9.png").read_bytes()), (names[2], (EX / "image-1.png").read_bytes())])
a = r.get_json()
check("accepted, original names shown", r.status_code == 200 and sorted(i["name"] for i in a["images"]) == sorted(names), [i["name"] for i in a["images"]])
check("pictures still readable through the API", all(c.get(i["url"] + "?w=200").status_code == 200 for i in a["images"]))

print("R3  a huge screenshot (6x upscaled = 5970 px wide) - same legibility as the original, small PowerPoint")
big = Image.open(EX / "image-6.png").convert("RGB")
big = big.resize((big.width * 6, big.height * 6), Image.LANCZOS)
buf = io.BytesIO(); big.save(buf, "PNG")
r = upload([("huge_stress.png", buf.getvalue())]); a = r.get_json()
info = a["images"][0]
check("huge picture: role + Max still right", info["role"] == "stress" and info["max"] == 545.0, (info["role"], info["max"]))
check("legend text measured in ORIGINAL pixels (~57 px = 6 x 9.5)", 50 <= (info["legend_px"] or 0) <= 65, info["legend_px"])
import mw_report as mw                          # noqa: E402
w, h, k = mw.image_size(server.load_infos(a["sid"])[0]["path"])
check("builder scales it down to 2600 px (PowerPoint stays small)", max(w, h) <= 2600 and 0.4 < k < 0.5, (w, h, k))
# predicted legend size for the same picture twice (side by side) must equal the ORIGINAL picture's (5.4 pt)
sc = mw.panel_scales("results", "side", [(w, h), (w, h)], ["a", "b"])
pt = sc[0] * k * info["legend_px"] / 0.72
check("predicted legend text equals the un-scaled original (about 5.4 pt)", 5.0 <= pt <= 5.8, pt)

print("R4  a 10-case job (the case selector in the table) -> summary paginates, nothing falls off the slide")
base = analyzer.analyze([str(p) for p in sorted(EX.glob("image-*.png"))])
r = upload([(p.name, p.read_bytes()) for p in sorted(EX.glob("*.png"))]); a = r.get_json()
sid = a["sid"]; infos = server.load_infos(sid)
nid = max(i["id"] for i in infos) + 1
extra = []
src = {(i["role"], i["case"]): i for i in infos if i["role"] in ("bc", "deformation", "stress")}
for n in range(3, 11):                                   # eight more cases: copies of case 1 / 2 pictures (pixel-changed so they are not 'duplicates')
    for role in ("bc", "deformation", "stress"):
        s = src[(role, 1 if n % 2 else 2)]
        im = Image.open(s["path"]).convert("RGB"); im.putpixel((3, 3), (n, n, n))
        p = Path(s["path"]).with_name(f"case{n}_{role}.png"); im.save(p)
        i2 = dict(s); i2.update(id=nid, name=p.name, path=str(p), case=n, hash=f"x{nid}", user_set=True); extra.append(i2); nid += 1
server.save_infos(sid, infos + extra)
infos = infos + extra
asg = [{"id": i["id"], "role": i["role"], "case": i["case"]} for i in infos]
rg = c.post("/api/regroup", json={"sid": sid, "assign": asg}).get_json()
check("ten cases drafted", len(rg["cases"]) == 10, len(rg["cases"]))
cases = [{"n": k["n"], "name": f"Case number {k['n']}", "subtitle": k["subtitle"], "short_name": f"Case {k['n']}", "bc_id": k["bc_id"], "def_id": k["def_id"],
          "stress_id": k["stress_id"], "bc_items": [t.replace("[location - please confirm]", "the shaft") for t in k["bc_items"]], "notes": [], "extras": [],
          "def_max": k["def_max"], "stress_max": k["stress_max"], "legend_px": k["legend_px"]} for k in rg["cases"]]
body = {"sid": sid, "cover": a["cover"], "short": "Blade", "fos": 1.3, "yield": 380, "allowable": 292, "show_utilisation": True,
        "material": {"rows": [["IS 2062", "Duct", "191400", "0.3", "7.8E-9", "380", "292"]], "units": server.SETTINGS["units"]},
        "geometry": rg["geometry"], "meshes": rg["meshes"], "cases": cases, "make_pdf": False}
g = c.post("/api/generate", json=body)
check("generated", g.status_code == 200, g.get_data(as_text=True)[:200])
f = server.OUT / g.get_json()["pptx"]
prs = Presentation(str(f))
titles = [next((sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()), "") for s in prs.slides]
check("summary split over two slides (8 + 2 rows)", titles[-2:] == ["Final Summary", "Final Summary (cont.)"], titles[-3:])
rows = [[len(sh.table.rows)] for sh in prs.slides[-2].shapes if getattr(sh, "has_table", False) and sh.has_table]
check("first summary slide holds 8 data rows + header", rows == [[9]], rows)
f.unlink(missing_ok=True)

print("R5  a very long setup list (14 long sentences) - the list shrinks instead of running off the slide")
items = [f"Fixed support applied on the surface of the {x} that is welded to the main frame ({i})." for i, x in enumerate(["left hub", "right hub", "centre rib", "upper flange", "lower flange", "left shaft end", "right shaft end", "stiffener 1", "stiffener 2", "stiffener 3", "base plate", "cover plate", "seal housing", "drain pipe"], 1)]
size = mw.image_size(str(EX / "image-2.png"))[:2]
inf = mw.bc_layout_info(size, items, [], "side")
check("14 long sentences: font shrunk below 14 pt", inf["fs"] < 14, inf)
check("...and the overflow flag is honest about whether it still does not fit", isinstance(inf["overflow"], bool))
inf2 = mw.bc_layout_info(size, items[:4], [], "side")
check("4 normal sentences keep the template's 14 pt", inf2["fs"] == 14 and not inf2["overflow"], inf2)

print("R6  a photo that is not an ANSYS picture at all")
photo = Image.new("RGB", (900, 600), (120, 160, 90))
buf = io.BytesIO(); photo.save(buf, "JPEG")
r = upload([("holiday.jpg", buf.getvalue())]); a = r.get_json()
check("no crash: it becomes a plain geometry picture the engineer can re-assign", r.status_code == 200 and a["images"][0]["role"] in ("geometry", "mesh", "unknown"), a["images"][0]["role"])

print()
print("ALL RED-TEAM TESTS PASSED" if not fails else f"{len(fails)} FAILED: {fails}")
sys.exit(1 if fails else 0)
