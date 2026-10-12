#!/usr/bin/env python3
"""
End-to-end regression test (no browser, no running server needed):
  1. upload the 9 example pictures in a shuffled order,
  2. fill in the things only a human knows (support locations, yield stress, client ...),
  3. generate the PowerPoint,
  4. compare it shape-by-shape with the reference deck (MLD_Blade_FEA_Report.pptx, made by hand earlier).

    python tests/e2e_blade.py [path/to/reference.pptx]
"""
import hashlib
import io
import json
import random
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "app"))
sys.path.insert(0, str(ROOT / "engine"))
import server                                   # noqa: E402
from pptx import Presentation                   # noqa: E402

REF = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("/home/user/MLD_Blade_FEA_Report.pptx")

# what the engineer types in the review page (identical to the hand-made deck)
LOCATIONS = {  # case -> {line index: full sentence}
    1: {2: "Fixed support applied on the left shaft.", 3: "Fixed support applied on the right shaft."},
    2: {0: "Cylindrical support applied on the left shaft.", 3: "Moment of 1,000,000 N\u00b7mm (1000 N\u00b7m) applied at the left shaft end.",
        4: "Cylindrical support applied on the right shaft.", 5: "Fixed support applied at the right shaft end."},
}


def dump(path):
    out = []
    for si, s in enumerate(Presentation(str(path)).slides, 1):
        for sh in s.shapes:
            txt = ""
            if sh.has_text_frame:
                txt = sh.text_frame.text
            elif getattr(sh, "has_table", False) and sh.has_table:
                txt = "|".join(c.text for r in sh.table.rows for c in r.cells)
            img = hashlib.md5(sh.image.blob).hexdigest()[:8] if sh.shape_type == 13 else ""
            out.append((si, sh.name, round(sh.left / 12700, 1), round(sh.top / 12700, 1), round(sh.width / 12700, 1),
                        round(sh.height / 12700, 1), txt, img))
    return out


def main():
    c = server.app.test_client()
    files = sorted((ROOT / "examples").glob("*.png"))
    random.Random(11).shuffle(files)
    r = c.post("/api/analyze", data={"files": [(io.BytesIO(f.read_bytes()), f.name) for f in files]},
               content_type="multipart/form-data")
    assert r.status_code == 200, r.get_data(as_text=True)
    a = r.get_json()
    print(f"analysed {len(a['images'])} pictures in {a['took']} s; warnings: {[w['text'] for w in a['warnings']] or 'none'}")

    cover = a["cover"]
    assert cover["title"] == "Static Structural Analysis Of MLD Blade", cover
    assert cover["report_no"] == "MWI/FEA/ACK-6203/01" and cover["client"] == "KCP" and cover["short"] == "Blade", cover
    mat = server.SETTINGS["material"]
    fos, yld = 1.3, 380
    allow = int(yld // fos)
    assert allow == 292
    rows = [[mat["material"], mat["item"], mat["E"], mat["nu"], mat["rho"], str(yld), str(allow)]]

    # live wording
    d = c.post("/api/derive", json={"allowable": allow, "yield": yld, "fos": fos, "short": cover["short"], "cases": [
        {"n": k["n"], "def_max": k["def_max"], "stress_max": k["stress_max"], "red_px": k["stress_red_px"],
         "legend": k["stress_legend"], "stress_unit": k["stress_unit"], "def_unit": k["def_unit"]} for k in a["cases"]]}).get_json()
    print("derive checks:")
    for ck in d["checks"]:
        print("   ", ck["level"].upper().ljust(4), ck["text"])
    print("   suggested yield from the legends:", d["suggest"])
    assert d["suggest"]["yield"] == 380

    cases = []
    for k in a["cases"]:
        items = list(k["bc_items"])
        for idx, text in LOCATIONS[k["n"]].items():
            items[idx] = text
        cases.append({"n": k["n"], "name": k["name"], "subtitle": k["subtitle"], "short_name": k["short_name"],
                      "bc_id": k["bc_id"], "def_id": k["def_id"], "stress_id": k["stress_id"], "bc_items": items,
                      "notes": [], "extras": [], "def_max": k["def_max"], "stress_max": k["stress_max"], "legend_px": k["legend_px"]})
    payload = {"sid": a["sid"], "cover": cover, "short": cover["short"], "fos": fos, "yield": yld, "allowable": allow, "show_utilisation": False,
               "material": {"rows": rows, "units": server.SETTINGS["units"]}, "geometry": a["geometry"],
               "meshes": a["meshes"], "cases": cases, "make_pdf": False}
    r = c.post("/api/generate", json=payload)
    assert r.status_code == 200, r.get_data(as_text=True)
    g = r.get_json()
    new = server.OUT / g["pptx"]
    print("generated:", new, f"({g['slides']} slides)")

    A, B = dump(REF), dump(new)
    # Expected, intended difference: Case 2's stress (545 MPa) is above the yield strength (380 MPa), so since v2 the app adds ONE
    # sentence saying the linear-elastic result is only indicative there.  Remove exactly that sentence, then everything must match.
    removed = 0
    B2 = []
    for row in B:
        lines = row[6].split("\n")
        keep = [l for l in lines if not l.startswith("The peak stress (")]
        removed += len(lines) - len(keep)
        B2.append(row[:6] + ("\n".join(keep),) + row[7:])
    assert removed == 1, f"expected exactly one 'above the yield strength' sentence, found {removed}"
    print("v2 adds exactly one engineering sentence (peak stress above yield) - removed for the comparison")
    B = B2

    def no_date(rows):                       # the cover date is 'today': it differs from the day the reference deck was made
        return [r[:6] + ("\n".join("Date: <today>" if l.startswith("Date: ") else l for l in r[6].split("\n")),) + r[7:] for r in rows]
    A, B = no_date(A), no_date(B)
    print(f"reference shapes: {len(A)}   generated shapes: {len(B)}   (cover date ignored: it is always today)")
    diffs = [(x, y) for x, y in zip(A, B) if x != y]
    for x, y in diffs[:12]:
        print("  DIFF\n    ref:", x, "\n    new:", y)
    ok = len(A) == len(B) and not diffs
    print("RESULT:", "IDENTICAL to the hand-made deck" if ok else f"{len(diffs)} differences")
    new.unlink(missing_ok=True)
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
