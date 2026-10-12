"""Browser test of version 2.2 (Playwright + Chromium).

    needs:  a running server  python app/server.py --host 0.0.0.0 --port 5055 --no-browser   (FEA_URL points to it)
    run:    CHROMIUM=/usr/bin/chromium FEA_URL=http://127.0.0.1:5055 python tests/browser_ui_v22.py

  A. the page shows its version
  B. only SOME pictures are uploaded: preview appears by itself, both downloads work, no error
  C. a slow reading shows a running seconds counter and can be cancelled
  D. the page inside a sandboxed frame (like the preview window of a website): downloads allowed / blocked
  E. 'Open the Downloads folder' button (program running on the user's own computer)
"""
import http.server
import os
import socketserver
import subprocess
import sys
import tempfile
import threading
import time
import urllib.parse
import urllib.request
from pathlib import Path

import pymupdf
from playwright.sync_api import sync_playwright
from pptx import Presentation

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
EX = ROOT / "examples"
SHOTS = HERE / "_shots"
SHOTS.mkdir(exist_ok=True)
URL = os.environ.get("FEA_URL", "http://127.0.0.1:5055")
LAUNCH = {"executable_path": os.environ["CHROMIUM"], "args": ["--no-sandbox"]} if os.environ.get("CHROMIUM") else {}
fails, errors = [], []


def check(name, cond, info=""):
    print(("  ok   " if cond else "  FAIL ") + name + (f"   [{info}]" if info and not cond else ""))
    if not cond:
        fails.append(name)


def pics(*nums):
    return [str(EX / f"image-{n}.png") for n in nums]


def wait_review(scope, secs=120):
    scope.wait_for_selector("#s-review:not([hidden])", timeout=secs * 1000)
    scope.wait_for_function("document.querySelector('#pv-img') && document.querySelector('#pv-img').naturalWidth > 100", timeout=60000)


class Wrap(http.server.BaseHTTPRequestHandler):
    """serves a parent page that shows the program inside an <iframe sandbox=...>"""
    def log_message(self, *a):
        pass

    def do_GET(self):
        q = urllib.parse.parse_qs(urllib.parse.urlparse(self.path).query)
        flags = q.get("f", [""])[0]
        body = (f'<!doctype html><body style="margin:0"><iframe src="{URL}/" sandbox="{flags}" '
                f'style="width:1280px;height:880px;border:0"></iframe></body>').encode()
        self.send_response(200); self.send_header("Content-Type", "text/html"); self.end_headers(); self.wfile.write(body)


class ReuseServer(socketserver.TCPServer):
    allow_reuse_address = True


def wait_server(port, secs=40):
    t0 = time.time()
    while time.time() - t0 < secs:
        try:
            with urllib.request.urlopen(f"http://127.0.0.1:{port}/api/status", timeout=2):
                return True
        except Exception:
            time.sleep(0.5)
    return False


with sync_playwright() as p:
    b = p.chromium.launch(**LAUNCH)
    ctx = b.new_context(viewport={"width": 1280, "height": 900}, accept_downloads=True)
    page = ctx.new_page()
    console = []
    page.on("console", lambda m: console.append((m.type, m.text)))
    page.on("pageerror", lambda e: errors.append(str(e)))
    dialogs = []
    page.on("dialog", lambda d: (dialogs.append(d.message), d.dismiss()))

    # ───────────────────────────────────────────── A
    print("A. the page shows its version")
    page.goto(URL)
    page.wait_for_selector("#chip-version")
    ver = page.locator("#chip-version").inner_text()
    check("version chip says v2.2", ver == "v2.2", ver)
    srcs = page.evaluate("Array.from(document.querySelectorAll('script[src],link[rel=stylesheet]')).map(e => e.src || e.href)")
    check("the page loads app.js and style.css with ?v=2.2", all("?v=2.2" in s for s in srcs) and len(srcs) == 2, str(srcs))

    # ───────────────────────────────────────────── B
    print("B. only SOME pictures: Case 1 (setup, deformation, stress) - no geometry, no mesh, no Case 2")
    page.set_input_files("#file", pics(2, 1, 3))
    wait_review(page)
    n_slides = int(page.locator("#pv-count").inner_text().split()[-1])
    print(f"     preview shows {n_slides} slides")
    check("review page and preview appear by themselves", page.locator("#pv-img").evaluate("e => e.naturalWidth") > 500)
    check("no error box is shown", page.locator("#alert").is_hidden())
    check("the report is not empty (cover, material, a case, summary ...)", n_slides >= 5, str(n_slides))
    page.screenshot(path=str(SHOTS / "v22_subset_review.png"))
    with page.expect_download(timeout=60000) as dl:
        page.locator("#btn-pptx").click()
    d = dl.value
    f = Path(tempfile.gettempdir()) / ("v22_" + d.suggested_filename)
    d.save_as(f)
    prs = Presentation(str(f))
    check("PowerPoint downloads and opens", len(prs.slides) == n_slides, f"{len(prs.slides)} vs {n_slides}")
    check("file name is FEA_Report_...pptx", d.suggested_filename.startswith("FEA_Report_") and d.suggested_filename.endswith(".pptx"), d.suggested_filename)
    toast = page.locator("#toasts .toast.dl").inner_text()
    check("toast: download started, saved in Downloads", "Download started" in toast and "Downloads folder" in toast, toast)
    check("toast on the program's own page has NO 'inside another window' note", "inside another window" not in toast)
    with page.expect_download(timeout=90000) as dl2:
        page.locator("#btn-pdf").click()
    d2 = dl2.value
    f2 = Path(tempfile.gettempdir()) / ("v22_" + d2.suggested_filename)
    d2.save_as(f2)
    check("PDF downloads with the same number of pages", len(pymupdf.open(str(f2))) == n_slides)
    check("no alert / confirm dialog", not dialogs)

    # ───────────────────────────────────────────── C
    print("C. a slow reading shows a running counter and can be cancelled")
    page.goto(URL)
    page.wait_for_selector("#chip-version")
    held = []
    page.route("**/api/analyze", lambda route: held.append(route))          # the server 'does not answer' for now
    page.set_input_files("#file", pics(1, 2, 3))
    page.wait_for_timeout(5600)
    check("the busy screen is up", page.locator("#busy").is_visible())
    extra = page.locator("#busy-extra").inner_text()
    check("it counts the seconds (so it never looks frozen)", extra.strip().endswith(" s") or " s" in extra, extra)
    check("a Cancel button is there", page.locator("#busy-cancel").is_visible())
    page.screenshot(path=str(SHOTS / "v22_busy.png"))
    page.locator("#busy-cancel").click()
    page.wait_for_timeout(500)
    check("Cancel closes the busy screen", page.locator("#busy").is_hidden())
    check("the upload page is usable again", page.locator("#s-upload").is_visible() and page.locator("#drop").is_visible())
    alert = page.locator("#alert").inner_text()
    check("a plain 'Cancelled' message (not 'Cannot reach the program')", "Cancelled" in alert and "Cannot reach" not in alert, alert)
    for r in held:
        try:
            r.abort()
        except Exception:
            pass
    page.unroute("**/api/analyze")
    page.set_input_files("#file", pics(1, 2, 3))                     # and a new try works
    wait_review(page)
    check("a new try after Cancel works", page.locator("#s-review").is_visible())

    # ───────────────────────────────────────────── D
    print("D. the page shown inside a sandboxed frame (preview window of a website)")
    wrap = ReuseServer(("127.0.0.1", 5090), Wrap)
    threading.Thread(target=wrap.serve_forever, daemon=True).start()
    for label, flags, expect_file in (
        ("frame WITH allow-downloads", "allow-scripts allow-same-origin allow-forms allow-popups allow-modals allow-downloads", True),
        ("frame WITHOUT allow-downloads", "allow-scripts allow-same-origin allow-forms allow-popups allow-modals", False),
    ):
        print("  -", label)
        pg = ctx.new_page()
        flog = []
        pg.on("console", lambda m, flog=flog: flog.append(m.text))
        pg.goto("http://localhost:5090/?f=" + urllib.parse.quote(flags))        # other origin than the program (127.0.0.1)
        fr = pg.frame_locator("iframe")
        fr.locator("#chip-version").wait_for(timeout=30000)
        check(f"{label}: the app works inside the frame (version chip)", fr.locator("#chip-version").inner_text() == "v2.2")
        fr.locator("#file").set_input_files(pics(1, 2, 3, 4, 5, 6, 7, 8, 9))
        fr.locator("#s-review:not([hidden])").wait_for(timeout=120000)
        fr.locator("#pv-img").wait_for(timeout=60000)
        time.sleep(1.5)
        got = None
        try:
            with pg.expect_download(timeout=9000) as dl3:
                fr.locator("#btn-pptx").click()
            got = dl3.value
        except Exception:
            got = None
        check(f"{label}: file " + ("arrives" if expect_file else "is blocked by the frame (as browsers do) - nothing arrives"), (got is not None) == expect_file)
        if got is not None:
            ff = Path(tempfile.gettempdir()) / ("v22f_" + got.suggested_filename)
            got.save_as(ff)
            check(f"{label}: the downloaded file is a real PowerPoint", len(Presentation(str(ff)).slides) == 11)
        else:
            check(f"{label}: the browser really blocked it (console says so)", any("Download is disallowed" in t or "allow-downloads" in t for t in flog), str(flog)[:200])
        t = fr.locator("#toasts .toast.dl")
        t.wait_for(timeout=10000)
        txt = t.inner_text()
        check(f"{label}: the message tells where the copy is, in case the window blocks", "inside another window" in txt and "outputs" in txt, txt)
        pg.screenshot(path=str(SHOTS / ("v22_frame_" + ("ok" if expect_file else "blocked") + ".png")))
        pg.close()

    wrap.shutdown(); wrap.server_close()

    # ───────────────────────────────────────────── E
    print("E. 'Open the Downloads folder' (program on the user's own computer)")
    logf = Path(tempfile.gettempdir()) / "fea_open_ui.log"
    logf.unlink(missing_ok=True)
    env = dict(os.environ, FEA_TEST_OPEN_LOG=str(logf))
    sp = subprocess.Popen([sys.executable, str(ROOT / "app" / "server.py"), "--port", "5056", "--no-browser"], cwd=ROOT, env=env,
                          stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    try:
        check("the local-mode test server started", wait_server(5056))
        pe = ctx.new_page()
        pe.on("pageerror", lambda e: errors.append(str(e)))
        pe.goto("http://127.0.0.1:5056/")
        pe.wait_for_selector("#chip-version")
        pe.set_input_files("#file", pics(1, 2, 3))
        wait_review(pe)
        with pe.expect_download(timeout=60000):
            pe.locator("#btn-pptx").click()
        pe.wait_for_selector("#btn-open-dl", timeout=10000)
        check("the toast has the button 'Open the Downloads folder'", pe.locator("#btn-open-dl").is_visible())
        pe.screenshot(path=str(SHOTS / "v22_open_downloads.png"))
        pe.locator("#btn-open-dl").click()
        t0 = time.time()
        while time.time() - t0 < 6 and not logf.exists():
            time.sleep(0.2)
        check("clicking it asks the computer to open the folder", logf.exists())
        pe.close()
    finally:
        sp.terminate()
        try:
            sp.wait(8)
        except subprocess.TimeoutExpired:
            sp.kill()

    check("no JavaScript errors in any page", not errors, str(errors)[:300])
    b.close()

print()
print("FAILS: " + (", ".join(fails) if fails else "none"))
sys.exit(1 if fails else 0)
